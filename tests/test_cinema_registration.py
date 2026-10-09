import unittest
from types import SimpleNamespace
from unittest.mock import patch

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import handlers.cinema_registration as cinema
import handlers.register_hello as hello
from db import Base, Event, Registration, User


class FakeBot:
    def __init__(self):
        self.next_steps = {}
        self.messages = []
        self.photos = []
        self.documents = []
        self.callback_answers = []
        self.message_handlers = {}
        self.callback_handlers = []

    def message_handler(self, **kwargs):
        def register(callback):
            self.message_handlers[tuple(kwargs.get("commands", []))] = callback
            return callback

        return register

    def callback_query_handler(self, **kwargs):
        def register(callback):
            self.callback_handlers.append((kwargs.get("func"), callback))
            return callback

        return register

    def send_message(self, chat_id, text, **kwargs):
        self.messages.append((chat_id, text, kwargs))

    def send_photo(self, chat_id, photo, **kwargs):
        self.photos.append((chat_id, photo, kwargs))

    def send_document(self, chat_id, document, **kwargs):
        self.documents.append((chat_id, document, kwargs))

    def answer_callback_query(self, call_id, *args, **kwargs):
        self.callback_answers.append((call_id, args, kwargs))

    def reply_to(self, message, text, **kwargs):
        self.messages.append((message.chat.id, text, kwargs))

    def register_next_step_handler(self, message, callback, *args):
        self.next_steps[message.chat.id] = (callback, args)

    def reply(self, chat_id, user_id, text=None, photo=None):
        message = SimpleNamespace(
            chat=SimpleNamespace(id=chat_id, type="private"),
            from_user=SimpleNamespace(id=user_id),
            text=text,
            photo=photo,
        )
        callback, args = self.next_steps.pop(chat_id)
        callback(message, *args)

    def dispatch_callback(self, call):
        for matches, callback in self.callback_handlers:
            if matches is None or matches(call):
                callback(call)
                return
        raise AssertionError(f"No callback handler matched {call.data!r}")


class CinemaRegistrationTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.session_factory = sessionmaker(
            bind=self.engine,
            expire_on_commit=False,
        )
        self.session_patch = patch.object(
            cinema,
            "SessionLocal",
            self.session_factory,
        )
        self.session_patch.start()

        with self.session_factory() as session:
            session.add(
                Event(
                    code="cinema",
                    name="Cinema",
                    is_active=True,
                    student_price=70000,
                    guest_price=80000,
                )
            )
            session.commit()

        self.bot = FakeBot()

    def tearDown(self):
        self.session_patch.stop()
        self.engine.dispose()

    def start(self, chat_id, registration_type="student", chat_type="private"):
        message = SimpleNamespace(
            chat=SimpleNamespace(id=chat_id, type=chat_type)
        )
        cinema.begin_registration(
            self.bot,
            message,
            chat_id,
            registration_type,
        )

    def reply(self, chat_id, text=None, photo=None):
        self.bot.reply(chat_id, chat_id, text=text, photo=photo)

    def registration(self, chat_id):
        with self.session_factory() as session:
            user = session.query(User).filter_by(tg_id=str(chat_id)).one()
            return (
                session.query(Registration)
                .filter_by(user_id=user.id)
                .one()
            )

    def add_approved_student_referrer(self):
        with self.session_factory() as session:
            event = session.query(Event).filter_by(code="cinema").one()
            user = User(
                tg_id="9000",
                full_name="Approved Student",
                student_id_card="54321",
                is_student=True,
            )
            session.add(user)
            session.flush()
            session.add(
                Registration(
                    user_id=user.id,
                    event_id=event.id,
                    registration_type="student",
                    status="approved",
                )
            )
            session.commit()

    def test_exit_cancels_each_student_step_and_allows_restart(self):
        cancellation_inputs = [
            (101, [], "exit"),
            (102, ["Student Name"], "/cancel"),
            (103, ["Student Name", "12345"], "EXIT"),
            (104, ["/start"], None),
        ]

        for chat_id, answers, cancel_input in cancellation_inputs:
            with self.subTest(chat_id=chat_id):
                self.start(chat_id)
                for answer in answers:
                    self.reply(chat_id, text=answer)
                if cancel_input is not None:
                    self.reply(chat_id, text=cancel_input)

                registration = self.registration(chat_id)
                self.assertEqual(registration.status, "cancelled")
                self.assertIsNone(registration.payment_proof_file_id)
                self.assertIsNone(registration.linked_student_id)

                self.start(chat_id)
                self.assertEqual(
                    self.registration(chat_id).status,
                    "collecting",
                )

    def test_exit_cancels_guest_steps_and_clears_referrer(self):
        self.add_approved_student_referrer()
        cancellation_inputs = [
            (201, [], "لغو"),
            (202, ["Guest Name"], "خروج"),
            (203, ["Guest Name", "54321"], "exit"),
        ]

        for chat_id, answers, cancel_input in cancellation_inputs:
            with self.subTest(chat_id=chat_id):
                self.start(chat_id, registration_type="guest")
                for answer in answers:
                    self.reply(chat_id, text=answer)
                self.reply(chat_id, text=cancel_input)

                registration = self.registration(chat_id)
                self.assertEqual(registration.status, "cancelled")
                self.assertIsNone(registration.linked_student_id)
                self.assertIsNone(registration.payment_proof_file_id)

    def test_invalid_replies_retry_and_only_approved_student_can_refer(self):
        self.start(301)
        self.reply(301, text="/help")
        self.assertIn(301, self.bot.next_steps)
        self.reply(301, text="Student Name")
        self.reply(301, text="not-a-number")
        self.assertIn(301, self.bot.next_steps)
        self.reply(301, text="۱۲۳۴۵")
        self.reply(301, photo=[SimpleNamespace(file_id="proof-301")])

        registration = self.registration(301)
        self.assertEqual(registration.status, "waiting_admin")
        with self.session_factory() as session:
            user = session.query(User).filter_by(tg_id="301").one()
            self.assertEqual(user.student_id_card, "12345")

        self.start(302, registration_type="guest")
        self.reply(302, text="Guest Name")
        self.reply(302, text="12345")
        self.assertIn(302, self.bot.next_steps)

        self.add_approved_student_referrer()
        self.reply(302, text="54321")
        self.reply(302, photo=[SimpleNamespace(file_id="proof-302")])
        guest_registration = self.registration(302)
        self.assertEqual(guest_registration.status, "waiting_admin")
        self.assertEqual(guest_registration.linked_student_id, "54321")

    def test_unexpected_error_mid_flow_notifies_user_without_crashing(self):
        self.start(601)

        with patch.object(
            cinema,
            "_get_collecting_registration",
            side_effect=RuntimeError("boom"),
        ):
            # Must not raise: an unexpected error inside a step should be
            # caught, not propagate up into telebot's polling loop.
            self.reply(601, text="Some Name")

        last_message = self.bot.messages[-1]
        self.assertEqual(last_message[0], 601)
        self.assertIn("خطا", last_message[1])
        self.assertIn("/start", last_message[1])

        # The DB write never happened (it blew up before the commit),
        # so the registration is still sitting in "collecting" and the
        # user can recover cleanly via /start.
        registration = self.registration(601)
        self.assertEqual(registration.status, "collecting")

    def test_unexpected_error_while_exiting_notifies_user_without_crashing(self):
        self.start(602)
        self.reply(602, text="Some Name")

        with patch.object(
            cinema,
            "_get_collecting_registration",
            side_effect=RuntimeError("boom"),
        ):
            # Exiting mid-flow must not crash even if the cancel lookup
            # itself blows up.
            self.reply(602, text="exit")

        last_message = self.bot.messages[-1]
        self.assertEqual(last_message[0], 602)
        self.assertIn("خطا", last_message[1])

    def test_group_chat_is_rejected_without_creating_registration(self):
        self.start(401, chat_type="group")
        self.assertNotIn(401, self.bot.next_steps)
        with self.session_factory() as session:
            self.assertIsNone(
                session.query(User).filter_by(tg_id="401").first()
            )

    def test_admin_command_shows_menu_instead_of_listing_directly(self):
        admin_id = next(iter(hello.ADMIN_IDS))

        with patch.object(hello, "SessionLocal", self.session_factory):
            hello.register_hello(self.bot)
            admin_handler = self.bot.message_handlers[("admin",)]
            admin_handler(
                SimpleNamespace(
                    chat=SimpleNamespace(id=admin_id),
                    from_user=SimpleNamespace(id=admin_id),
                )
            )

        self.assertEqual(len(self.bot.messages), 1)
        chat_id, text, kwargs = self.bot.messages[-1]
        self.assertEqual(chat_id, admin_id)
        self.assertIn("پنل ادمین", text)

        markup = kwargs["reply_markup"]
        callbacks = {
            button.callback_data
            for row in markup.keyboard
            for button in row
        }
        self.assertEqual(callbacks, {"admin_menu:pending", "admin_menu:export"})

    def test_admin_menu_pending_option_lists_waiting_registrations(self):
        admin_id = next(iter(hello.ADMIN_IDS))

        with self.session_factory() as session:
            event = session.query(Event).filter_by(code="cinema").one()
            user = User(tg_id="701", full_name="Waiting User")
            session.add(user)
            session.flush()
            session.add(
                Registration(
                    user_id=user.id,
                    event_id=event.id,
                    registration_type="student",
                    status="waiting_admin",
                )
            )
            session.commit()

        with patch.object(hello, "SessionLocal", self.session_factory):
            hello.register_hello(self.bot)
            call = SimpleNamespace(
                id="cb-pending",
                data="admin_menu:pending",
                from_user=SimpleNamespace(id=admin_id),
                message=SimpleNamespace(chat=SimpleNamespace(id=admin_id)),
            )
            self.bot.dispatch_callback(call)

        self.assertEqual(len(self.bot.messages), 1)
        chat_id, text, kwargs = self.bot.messages[-1]
        self.assertEqual(chat_id, admin_id)
        self.assertIn("Waiting User", text)
        callbacks = {
            button.callback_data for row in kwargs["reply_markup"].keyboard for button in row
        }
        self.assertTrue(any(c.startswith("approve:") for c in callbacks))
        self.assertTrue(any(c.startswith("reject:") for c in callbacks))

    def test_admin_menu_export_rejects_non_admin(self):
        with patch.object(hello, "SessionLocal", self.session_factory):
            hello.register_hello(self.bot)
            call = SimpleNamespace(
                id="cb-export",
                data="admin_menu:export",
                from_user=SimpleNamespace(id=1),
                message=SimpleNamespace(chat=SimpleNamespace(id=1)),
            )
            self.bot.dispatch_callback(call)

        self.assertEqual(self.bot.documents, [])
        self.assertEqual(len(self.bot.callback_answers), 1)
        self.assertIn("اجازه", self.bot.callback_answers[0][1][0])

    def test_admin_export_excel_contains_approved_users_with_referrer(self):
        with self.session_factory() as session:
            event = session.query(Event).filter_by(code="cinema").one()

            student = User(
                tg_id="9001",
                full_name="Approved Student",
                student_id_card="11111",
                is_student=True,
            )
            session.add(student)
            session.flush()
            session.add(
                Registration(
                    user_id=student.id,
                    event_id=event.id,
                    registration_type="student",
                    status="approved",
                    ticket_code="CINEMA-11111",
                )
            )

            guest = User(
                tg_id="9002",
                full_name="Approved Guest",
                is_student=False,
            )
            session.add(guest)
            session.flush()
            session.add(
                Registration(
                    user_id=guest.id,
                    event_id=event.id,
                    registration_type="guest",
                    status="approved",
                    linked_student_id="11111",
                    ticket_code="CINEMA-22222",
                )
            )

            # A non-approved registration must not show up in the export.
            pending_user = User(tg_id="9003", full_name="Still Pending")
            session.add(pending_user)
            session.flush()
            session.add(
                Registration(
                    user_id=pending_user.id,
                    event_id=event.id,
                    registration_type="guest",
                    status="waiting_admin",
                )
            )
            session.commit()

        with patch.object(hello, "SessionLocal", self.session_factory):
            hello._send_approved_excel(self.bot, 999)

        self.assertEqual(len(self.bot.documents), 1)
        chat_id, document, kwargs = self.bot.documents[0]
        self.assertEqual(chat_id, 999)
        self.assertEqual(
            kwargs.get("visible_file_name"), "cinema_approved_users.xlsx"
        )

        from openpyxl import load_workbook

        document.seek(0)
        workbook = load_workbook(document)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))

        self.assertEqual(
            rows[0],
            (
                "ردیف",
                "نام و نام خانوادگی",
                "نوع ثبت‌نام",
                "آیدی تلگرام",
                "شماره دانشجویی",
                "آیدی معرف",
                "کد حضور",
            ),
        )
        # Only the two approved registrations show up (3 rows total: header
        # + student + guest), and each shows the right id columns:
        # students carry a student id with no referrer, guests carry a
        # referrer id with no student id.
        self.assertEqual(len(rows), 3)
        self.assertEqual(
            rows[1][:4],
            (1, "Approved Student", "دانشجو", "9001"),
        )
        self.assertEqual(rows[1][6], "CINEMA-11111")
        self.assertEqual(
            rows[2][:4],
            (2, "Approved Guest", "مهمان", "9002"),
        )
        self.assertEqual(rows[2][5], "11111")
        self.assertEqual(rows[2][6], "CINEMA-22222")

    def test_start_recovers_stale_registration_without_step_handler(self):
        with self.session_factory() as session:
            event = session.query(Event).filter_by(code="cinema").one()
            user = User(tg_id="501")
            session.add(user)
            session.flush()
            session.add(
                Registration(
                    user_id=user.id,
                    event_id=event.id,
                    registration_type="student",
                    status="collecting",
                )
            )
            session.commit()

        with patch.object(hello, "SessionLocal", self.session_factory):
            hello.register_hello(self.bot)
            start_handler = self.bot.message_handlers[("start",)]
            start_handler(
                SimpleNamespace(
                    chat=SimpleNamespace(id=501),
                    from_user=SimpleNamespace(
                        id=501,
                        first_name="Student",
                    ),
                )
            )

        message_text = self.bot.messages[-1][1]
        self.assertIn("نیمه‌تمام", message_text)
        self.assertEqual(
            self.bot.messages[-1][0],
            501,
        )
        markup = self.bot.messages[-1][2]["reply_markup"]
        self.assertEqual(
            markup.keyboard[0][0].callback_data,
            "register:student:cinema",
        )

    def test_capacity_cap_blocks_new_signups_once_full(self):
        with self.session_factory() as session:
            event = session.query(Event).filter_by(code="cinema").one()
            event.capacity = 2
            user = User(tg_id="9101")
            session.add(user)
            session.flush()
            session.add(
                Registration(
                    user_id=user.id,
                    event_id=event.id,
                    registration_type="student",
                    status="waiting_admin",
                )
            )
            session.commit()

        # 1 of 2 slots used so far - a new signup should proceed normally.
        self.start(201)
        self.assertIn(201, self.bot.next_steps)
        self.assertFalse(any("ظرفیت" in text for _, text, _ in self.bot.messages))

        with self.session_factory() as session:
            event = session.query(Event).filter_by(code="cinema").one()
            user = session.query(User).filter_by(tg_id="201").one()
            registration = (
                session.query(Registration)
                .filter_by(user_id=user.id, event_id=event.id)
                .one()
            )
            registration.status = "waiting_admin"
            session.commit()

        # Capacity (2) is now full - the next signup attempt must be blocked
        # with the "capacity full" message instead of starting the wizard.
        self.start(202)
        self.assertNotIn(202, self.bot.next_steps)
        self.assertIn("ظرفیت", self.bot.messages[-1][1])


if __name__ == "__main__":
    unittest.main()
